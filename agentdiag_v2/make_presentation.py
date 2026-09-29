"""Generate a fresh, result-aligned presentation inside AgentDiag v2."""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


NAVY = RGBColor(28, 45, 64)
TEAL = RGBColor(0, 116, 122)
GRAY = RGBColor(83, 92, 99)


def _text(slide, x, y, w, h, content, size=20, bold=False, color=NAVY):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()
    for i, line in enumerate(content.split("\n")):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(9)
    return box


def _slide(prs, title, body=None, foot=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _text(slide, 0.65, 0.35, 11.9, 0.65, title, 28, True)
    if body:
        _text(slide, 0.75, 1.25, 11.8, 5.25, body, 19)
    if foot:
        _text(slide, 0.75, 7.02, 11.8, 0.27, foot, 10, False, GRAY)
    return slide


def build(summary_path: Path, output_path: Path) -> None:
    summary = json.loads(summary_path.read_text())
    m = summary["test_calibrated"]
    prior = summary["test_position_prior"]
    temporal = summary["temporal"]
    keyword = summary["temporal_keyword_baseline"]
    count = summary["four_agent_counts"]
    missing = summary["missing_evidence"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    _slide(prs, "AgentDiag v2", "Probabilistic root-cause attribution for multi-agent LLM traces\n"
           "Fresh implementation and held-out evaluation\n"
           "Who&When benchmark", "Results generated from agentdiag_v2 artifacts")
    _slide(prs, "Problem and intended outputs",
           "Given a failed execution trace, identify the most likely responsible agent.\n"
           "Return a probability for every candidate agent; all probabilities sum to 1.\n"
           "Handle missing agent logs by marginalizing their evidence.\n"
           "Track reliability through turns and test whether warnings arrive before the first faulty message.",
           "Post-mortem diagnosis and early warning are evaluated separately")
    _slide(prs, "Data and evaluation protocol",
           f"184 traces parsed; question IDs kept together across partitions.\n"
           f"Primary four-agent experiment: train {count['train']}, validation {count['validation']}, test {count['test']}.\n"
           "Training fits probabilities; validation selects features, calibration, and warning threshold.\n"
           "The original BN was frozen before its test evaluation; later exploratory models inspected this test set.",
           "The source benchmark contains failed traces only")
    _slide(prs, "System architecture",
           "Trace JSON → validated agents and message order → evidence flags\n"
           "Evidence → categorical culprit Bayesian network → custom Variable Elimination\n"
           "Posterior → validation-fitted temperature scaling → root-cause distribution\n"
           "Turn evidence → HMM filtering → early-warning evaluation",
           "The network uses data-driven CPTs with smoothing")
    _slide(prs, "Inference and missing logs",
           "A single categorical variable represents the culprit, so exactly one root cause is modeled.\n"
           "Each agent contributes observable error, uncertainty, and correction signals.\n"
           "Custom Variable Elimination marginalizes missing evidence.\n"
           "Small-network results are cross-checked against full joint enumeration.",
           "A missing log is never treated as a clean log")
    _slide(prs, "Held-out attribution and calibration",
           f"Bayesian network: {m['correct']}/{m['n']} correct ({m['accuracy']:.1%}); "
           f"Brier {m['brier']:.3f}; ECE {m['ece']:.3f}.\n"
           f"Position-prior baseline: {prior['correct']}/{prior['n']} correct ({prior['accuracy']:.1%}); "
           f"Brier {prior['brier']:.3f}; ECE {prior['ece']:.3f}.\n"
           f"Accuracy bootstrap interval: {m['accuracy_95pct_bootstrap_interval'][0]:.1%}–"
           f"{m['accuracy_95pct_bootstrap_interval'][1]:.1%}.\n"
           "Small test set: competitive accuracy and superior calibration are not established.",
           "No direct AgenTracer claim without a matched protocol")
    slide = _slide(prs, "Reliability diagram", foot="Held-out test set; point size indicates number of traces per bin")
    image = summary_path.parent.parent / "phase5_calibration_evaluation" / "reliability_diagram.png"
    slide.shapes.add_picture(str(image), Inches(3.65), Inches(1.1), width=Inches(6.0))
    _slide(prs, "Missing evidence and temporal monitoring",
           f"Complete logs: {missing['complete']['accuracy']:.1%} attribution accuracy.\n"
           f"One random missing agent: {missing['one_random_missing']['accuracy']:.1%}; "
           f"two missing agents: {missing['two_random_missing']['accuracy']:.1%}.\n"
           f"HMM correct early warnings: {temporal['correct_early']}/{temporal['eligible']} eligible traces.\n"
           f"HMM wrong-agent warning traces: {temporal['false_agent_warning_traces']}/{temporal['total']}.\n"
           f"Keyword baseline: {keyword['correct_early']}/{keyword['eligible']} early; "
           f"{keyword['false_agent_warning_traces']}/{keyword['total']} wrong-agent warnings.",
           "Missing-log results use the same validation-fitted calibration")
    triage = summary["evidence_triage"]
    _slide(prs, "New evidence-backed review output",
           f"The true agent is in the BN's top-two shortlist for "
           f"{triage['top_two_contains_actual']}/{triage['test_count']} four-agent test traces "
           f"({triage['top_two_coverage']:.1%}).\n"
           "Each packet shows ranked probabilities and excerpts with source-message indices.\n"
           "Leave-one-agent-evidence-out sensitivity shows how the saved model reacts to that agent's evidence.\n"
           "This supports human review; top-choice accuracy remains 5/13, and the ablation is not a causal intervention.",
           "Phase 9: agentdiag_v2/artifacts/phase9_evidence_triage/")
    _slide(prs, "What is complete, and what remains unproven",
           "Implemented: parsing, data-driven BN, hand-written VE, valid probabilities, missing-log inference, "
           "calibration evaluation, evidence review packets, HMM, and reproducible phase outputs.\n"
           "Unproven: high top-choice accuracy, reliable early warning, publication-level novelty, and competitive results versus AgenTracer.\n"
           "Multimodal performance is not claimed: no separate modality-specific test was run.\n"
           "Use the measured results in this deck; improve evidence or add suitable labeled data before stronger claims.",
           "Full evidence: agentdiag_v2/artifacts/phase8_final_report/Final_Report.md")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)


if __name__ == "__main__":
    root = Path(__file__).resolve().parent / "artifacts" / "phase8_final_report"
    build(root / "summary.json", root / "AgentDiag_v2_Final_Review.pptx")
