#!/usr/bin/env python3
"""Independently recheck saved phase artifacts and headline metrics."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation

if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from agentdiag_v2.src.data import load_records, write_json
from agentdiag_v2.src.evaluation import metric_summary, normalized


HERE = Path(__file__).resolve().parent


def read_jsonl(path: Path) -> list[dict]:
    return [json.loads(line) for line in path.read_text().splitlines() if line.strip()]


def verify(artifacts: Path, dataset: Path) -> dict:
    checks = {}
    manifest = json.loads((artifacts / "phase_manifest.json").read_text())
    checks["all_ten_phase_reports_exist"] = (set(manifest) == {f"phase{i}" for i in range(10)}
                                              and all(Path(p).is_file() for p in manifest.values()))
    split_data = json.loads((artifacts / "phase0_protocol" / "splits.json").read_text())
    splits = split_data["splits"]
    ids = [trace_id for part in splits.values() for trace_id in part]
    records, exclusions = load_records(dataset)
    by_id = {r["id"]: r for r in records}
    checks["all_source_records_accounted_for"] = len(records) + len(exclusions) == 184 and set(ids) == set(by_id)
    checks["trace_splits_disjoint"] = len(ids) == len(set(ids))
    qid_to_part = {}
    group_safe = True
    for part, part_ids in splits.items():
        for trace_id in part_ids:
            qid = by_id[trace_id]["question_id"]
            if qid in qid_to_part and qid_to_part[qid] != part:
                group_safe = False
            qid_to_part[qid] = part
    checks["question_groups_do_not_leak"] = group_safe
    predictions = read_jsonl(artifacts / "phase5_calibration_evaluation" / "test_predictions.jsonl")
    exploratory = read_jsonl(artifacts / "phase5_calibration_evaluation" / "all_agent_counts_test_predictions.jsonl")
    test_ids = set(splits["test"])
    checks["test_predictions_only_test_traces"] = all(p["id"] in test_ids for p in predictions)
    checks["every_test_agent_count_evaluated"] = (set(p["id"] for p in exploratory) == test_ids
                                                   and all(normalized(p["probabilities"]) for p in exploratory))
    checks["all_test_probabilities_valid"] = all(normalized(p["probabilities"]) for p in predictions)
    reported = json.loads((artifacts / "phase5_calibration_evaluation" / "calibration.json").read_text())["calibrated_test"]
    recomputed = metric_summary(predictions)
    checks["headline_metrics_recomputed"] = all(abs(recomputed[k] - reported[k]) < 1e-12
                                                 for k in ("accuracy", "brier", "ece", "log_loss"))
    missing = read_jsonl(artifacts / "phase4_missing_evidence" / "missing_evidence_predictions.jsonl")
    complete = {p["id"]: p for p in missing if p["scenario"] == "complete"}
    checks["missing_scenarios_valid"] = (all(normalized(p["probabilities"]) for p in missing)
                                            and len(missing) == 4 * len(predictions))
    checks["complete_mask_matches_calibrated_prediction"] = all(
        max(abs(a - b) for a, b in zip(p["probabilities"], complete[p["id"]]["probabilities"])) < 1e-12
        for p in predictions)
    packets = read_jsonl(artifacts / "phase9_evidence_triage" / "test_review_packets.jsonl")
    packet_by_id = {p["id"]: p for p in packets}
    checks["review_packets_match_predictions"] = (
        len(packet_by_id) == len(predictions)
        and all(p["id"] in packet_by_id
                and [c["agent"] for c in packet_by_id[p["id"]]["review"]["candidates"][:2]]
                == packet_by_id[p["id"]]["review"]["top_two_agents"]
                and all(abs(c["probability"] - p["probabilities"][by_id[p["id"]]["agents"].index(c["agent"])]) < 1e-12
                        for c in packet_by_id[p["id"]]["review"]["candidates"])
                for p in predictions))
    checks["review_excerpts_have_valid_source_indices"] = all(
        any(m["raw_index"] == observation["raw_message_index"] and m["agent"] == candidate["agent"]
            for m in by_id[packet_row["id"]]["messages"])
        for packet_row in packets for candidate in packet_row["review"]["candidates"]
        for observation in candidate["observations"] if "raw_message_index" in observation)
    triage = json.loads((artifacts / "phase9_evidence_triage" / "metrics.json").read_text())
    checks["review_top_two_metric_recomputed"] = (
        triage["top_two_contains_actual"] == sum(p["top_two_contains_actual"] for p in packets)
        and triage["top_one_correct"] == sum(p["top_choice_correct"] for p in packets))
    warnings = read_jsonl(artifacts / "phase6_temporal_hmm" / "test_warnings.jsonl")
    checks["warning_labels_consistent"] = all(
        w["correct_early"] == (w["culprit_warning"] is not None
                                and w["culprit_warning"] < w["mistake_step"])
        and by_id[w["id"]]["label_message_matches"] for w in warnings)
    checks["reliability_figure_exists"] = (artifacts / "phase5_calibration_evaluation" / "reliability_diagram.png").is_file()
    checks["saved_model_and_calibrator_exist"] = (
        (artifacts / "phase3_bayesian_inference" / "model.json").is_file()
        and (artifacts / "phase5_calibration_evaluation" / "calibration.json").is_file()
        and (artifacts / "phase6_temporal_hmm" / "hmm_model.json").is_file())
    ppt_path = artifacts / "phase8_final_report" / "AgentDiag_v2_Final_Review.pptx"
    if ppt_path.is_file():
        presentation = Presentation(ppt_path)
        all_text = " ".join(shape.text for slide in presentation.slides for shape in slide.shapes
                            if shape.has_text_frame)
        checks["final_presentation_contains_measured_results"] = (
            len(presentation.slides) == 10
            and f"{reported['correct']}/{reported['n']}" in all_text
            and "not established" in all_text)
    else:
        checks["final_presentation_contains_measured_results"] = False
    outcome = {"passed": all(checks.values()), "checks": checks,
               "test_predictions": len(predictions), "temporal_test_traces": len(warnings)}
    phase8 = artifacts / "phase8_final_report"
    write_json(phase8 / "verification.json", outcome)
    (phase8 / "Verification_Report.md").write_text(
        "# AgentDiag v2 verification\n\n"
        + "\n".join(f"- {'PASS' if result else 'FAIL'} — {name.replace('_', ' ')}"
                    for name, result in checks.items())
        + f"\n\nOverall: {'PASS' if outcome['passed'] else 'FAIL'}. "
          f"Checked {len(predictions)} attribution predictions and {len(warnings)} temporal test traces.\n")
    return outcome


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--artifacts", type=Path, default=HERE / "artifacts")
    parser.add_argument("--dataset", type=Path, default=HERE.parent / "Agents_Failure_Attribution" / "Who&When")
    args = parser.parse_args()
    result = verify(args.artifacts.resolve(), args.dataset.resolve())
    print("PASS" if result["passed"] else "FAIL", "—", len(result["checks"]), "artifact checks")
    if not result["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
